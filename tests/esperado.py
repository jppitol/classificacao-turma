"""Valores esperados, calculados em Python direto da planilha (caminho independente do SQL)."""
import json, sys, warnings
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as ci
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from estrutura import SEMESTRES, PESOS_NCF
warnings.filterwarnings("ignore")
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
of = wb["Classificação OFICIAL"]; pf = wb["Preferências "]

seg, alunos = "ativa", {}
for r in range(3, of.max_row + 1):
    v = of.cell(r, 1).value
    if isinstance(v, str) and v.strip().upper() == "NOMES": seg = "reserva"; continue
    if isinstance(v, str) and isinstance(of.cell(r, 6).value, (int, float)): alunos[v.strip().upper()] = seg

def linhas(ws):
    return {ws.cell(r,1).value.strip().upper(): r for r in range(3, ws.max_row+1)
            if isinstance(ws.cell(r,1).value, str) and ws.cell(r,1).value.strip().upper() in alunos}

# 1) Semestre 1 isolado, para alunos com TODAS as notas do 1º lançadas: comparar com AK (MG) e AR (MM) da planilha
ws1 = wb["1º Semestre"]; L1 = linhas(ws1)
sem1 = {}
for nome, r in L1.items():
    cols = [c for comp in ("graduacao","militar") for _,_,avs in SEMESTRES[1][comp] for _,p,c in avs if p > 0 and c != "AQ"]
    if all(isinstance(ws1.cell(r, ci(c)).value,(int,float)) and ws1.cell(r, ci(c)).value > 0 for c in cols):
        sem1[nome] = {"mg": ws1[f"AK{r}"].value, "mm": ws1[f"AR{r}"].value}

# 1b) Média de cada disciplina do 1º semestre, como a planilha calcula (coluna de média de cada uma;
#     o que falta conta zero, ex.: recuperação = só a VF)
COL_MEDIA_SEM1 = {"Física I": "F", "Cálculo I": "K", "Álgebra Linear I": "P", "IC": "T", "Química I": "Y",
                  "Laboratório de Física I": "AD", "Laboratório de Química I": "AJ", "Instrução Militar I": "AP"}
disc1 = {nome: {d: ws1[f"{c}{r}"].value for d, c in COL_MEDIA_SEM1.items() if isinstance(ws1[f"{c}{r}"].value, (int, float))}
         for nome, r in L1.items()}

# 2) NCF atual, recalculado em Python com a regra da planilha: avaliação sem nota conta como zero
#    (divide pelo peso de TODAS as avaliações do curso, não só das lançadas)
W = {}  # (sem, comp) -> [(w, col)]
for n, cfg in SEMESTRES.items():
    for comp in ("graduacao","militar"):
        sd = sum(pd for _,pd,_ in cfg[comp])
        for _, pd, avs in cfg[comp]:
            sa = sum(p for _,p,_ in avs)
            for _, p, c in avs:
                W.setdefault(comp, []).append((n, (pd/sd)*(p/sa) if sd and sa else 0, c))
parcial = {}
L = {n: (wb[cfg["aba"]], linhas(wb[cfg["aba"]])) for n, cfg in SEMESTRES.items()}
for nome, sg in alunos.items():
    res = {}
    lancadas = 0
    for comp in ("graduacao","militar"):
        num = 0
        for n, w, c in W[comp]:
            ws, Ln = L[n]
            if nome not in Ln or w == 0: continue
            v = ws.cell(Ln[nome], ci(c)).value
            if isinstance(v,(int,float)) and 0 < v <= 10: num += w*v; lancadas += 1
        res[comp] = num / sum(w for _, w, _ in W[comp])
    pg, pm = PESOS_NCF[sg]
    if lancadas == 0: res = {"graduacao": None, "militar": None}
    g, m = res["graduacao"], res["militar"]
    ncf = pg*g + pm*m if lancadas else None
    parcial[nome] = {"segmento": sg, "mg": res["graduacao"], "mm": res["militar"], "ncf": ncf}

# 3) Distribuição da planilha (com o NCF que ela usou) para testar o algoritmo
alocacao = {}
for r in range(5, pf.max_row+1):
    nome = pf.cell(r,3).value
    if isinstance(nome,str) and nome.strip().upper() in alunos:
        alocacao[nome.strip().upper()] = pf.cell(r,4).value
ncf_planilha = {}
for r in range(3, 97):
    for cn, cv in ((9,10),(13,14)):
        nm, v = of.cell(r,cn).value, of.cell(r,cv).value
        if isinstance(nm,str) and isinstance(v,(int,float)): ncf_planilha[nm.strip().upper()] = v

json.dump({"sem1": sem1, "disc1": disc1, "parcial": parcial, "alocacao": alocacao, "ncf_planilha": ncf_planilha},
          open(Path(__file__).parent / "esperado.json", "w"), ensure_ascii=False, indent=1)
print(len(sem1), "alunos com 1º semestre completo;", len(parcial), "parciais;", len(alocacao), "alocações")
