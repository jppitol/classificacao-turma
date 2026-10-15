"""Gera os SQLs de carga a partir de estrutura.py e da planilha.

Uso:
    python scripts/gerar_sql.py                        # só a estrutura -> supabase/02_seed.sql
    python scripts/gerar_sql.py CLASSIFICACAO.xlsx     # + alunos, notas e preferências
                                                       #   -> supabase/03_importar_planilha.sql

O 03 contém as notas reais da turma: está no .gitignore e NUNCA deve ir pro GitHub.
"""
import sys
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from estrutura import ENGENHARIAS, PESOS_NCF, SEMESTRES, VAGAS  # noqa: E402

TURMA = "XXX"
SEP = ",\n"  # (fora das f-strings: compatível com Python 3.8+)
RAIZ = Path(__file__).resolve().parent.parent


def q(s):
    return "'" + str(s).replace("'", "''") + "'"


def turma_id():
    return f"(select id from public.turmas where nome = {q(TURMA)})"


def seed():
    out = ["-- Estrutura da Turma XXX (gerado por scripts/gerar_sql.py a partir de estrutura.py)",
           "begin;",
           f"insert into public.turmas (nome) values ({q(TURMA)});"]
    for seg, (pg, pm) in PESOS_NCF.items():
        out.append(f"insert into public.pesos_ncf values ({turma_id()}, {q(seg)}, {pg}, {pm});")
    for n in SEMESTRES:
        out.append(f"insert into public.semestres (turma_id, numero) values ({turma_id()}, {n});")
    for n, cfg in SEMESTRES.items():
        sem = f"(select id from public.semestres where turma_id = {turma_id()} and numero = {n})"
        for comp in ("graduacao", "militar"):
            for i, (nome, peso, avs) in enumerate(cfg[comp], 1):
                out.append(f"insert into public.disciplinas (semestre_id, nome, componente, peso, ordem) "
                           f"values ({sem}, {q(nome)}, {q(comp)}, {peso}, {i});")
                disc = f"(select id from public.disciplinas where semestre_id = {sem} and nome = {q(nome)})"
                vals = ", ".join(f"({disc}, {q(a)}, {p}, {j})" for j, (a, p, _) in enumerate(avs, 1))
                out.append(f"insert into public.avaliacoes (disciplina_id, nome, peso, ordem) values {vals};")
    out.append("insert into public.engenharias (sigla, nome) values "
               + ", ".join(f"({q(s)}, {q(n)})" for s, n in ENGENHARIAS) + ";")
    for seg, vs in VAGAS.items():
        out.append("insert into public.vagas values "
                   + ", ".join(f"({turma_id()}, {q(seg)}, {q(e)}, {k})" for e, k in vs.items()) + ";")
    out.append("commit;")
    return "\n".join(out) + "\n"


def importar(xlsx):
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # Lista e segmento: aba "Classificação OFICIAL" (ativa até a linha "NOMES", depois reserva)
    of = wb["Classificação OFICIAL"]
    alunos, seg = {}, "ativa"
    for r in range(3, of.max_row + 1):
        v = of.cell(r, 1).value
        if not isinstance(v, str) or not v.strip():
            continue
        if v.strip().upper() == "NOMES":
            seg = "reserva"
            continue
        if isinstance(of.cell(r, 6).value, (int, float)):  # só linhas com NCF calculado
            alunos[v.strip().upper()] = seg

    out = ["-- Alunos, notas e preferências importados da planilha.",
           "-- CONTÉM DADOS REAIS: não versionar, não publicar.",
           "begin;"]
    out.append("insert into public.alunos (turma_id, nome_guerra, segmento) values "
               + ",\n  ".join(f"({turma_id()}, {q(n)}, {q(s)})" for n, s in alunos.items()) + ";")

    # Notas: uma tabela de valores (nome, semestre, disciplina, avaliação, nota) ligada à estrutura por nome
    total, rows = 0, []
    for n, cfg in SEMESTRES.items():
        ws = wb[cfg["aba"]]
        linhas = {}
        for r in range(3, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and v.strip().upper() in alunos:
                linhas[v.strip().upper()] = r
        for comp in ("graduacao", "militar"):
            for disc, _, avs in cfg[comp]:
                for a, _, col in avs:
                    for nome, r in linhas.items():
                        v = ws.cell(r, ci(col)).value
                        # vazio ou 0 = não lançado (a planilha não distingue os dois)
                        if isinstance(v, (int, float)) and 0 < v <= 10:
                            rows.append(f"({q(nome)},{n},{q(disc)},{q(a)},{float(v)!r})")
    total = len(rows)
    out.append(f"""insert into public.notas (aluno_id, avaliacao_id, valor)
select al.id, a.id, v.valor
from (values
{SEP.join(rows)}
) as v(nome, semestre, disciplina, avaliacao, valor)
join public.alunos al      on al.turma_id = {turma_id()} and al.nome_guerra = v.nome
join public.semestres s    on s.turma_id = {turma_id()} and s.numero = v.semestre
join public.disciplinas d  on d.semestre_id = s.id and d.nome = v.disciplina
join public.avaliacoes a   on a.disciplina_id = d.id and a.nome = v.avaliacao;""")

    pf = wb["Preferências "]
    siglas = {s for s, _ in ENGENHARIAS}
    prefs, n_listas = [], 0
    for r in range(5, pf.max_row + 1):
        nome = pf.cell(r, 3).value
        if not isinstance(nome, str) or nome.strip().upper() not in alunos:
            continue
        ops = []
        for c in range(5, 15):
            o = pf.cell(r, c).value
            if o in siglas and o not in ops:
                ops.append(o)
        if len(ops) < 10:
            print(f"aviso: {nome} tem só {len(ops)} opção(ões) válida(s) distinta(s): {ops}", file=sys.stderr)
        if ops:
            n_listas += 1
        prefs += [f"({q(nome.strip().upper())},{q(e)},{i})" for i, e in enumerate(ops, 1)]
    if prefs:
        out.append(f"""insert into public.preferencias (aluno_id, engenharia, ordem)
select al.id, v.engenharia, v.ordem
from (values
{SEP.join(prefs)}
) as v(nome, engenharia, ordem)
join public.alunos al on al.turma_id = {turma_id()} and al.nome_guerra = v.nome;""")
    # conferência: se algum nome/disciplina não casou, desfaz tudo
    out.append(f"""do $$ begin
  if (select count(*) from public.alunos) <> {len(alunos)} then raise exception 'alunos: contagem errada'; end if;
  if (select count(*) from public.notas) <> {total} then raise exception 'notas: alguma linha não casou com a estrutura'; end if;
  if (select count(*) from public.preferencias) <> {len(prefs)} then raise exception 'preferências: contagem errada'; end if;
end $$;""")
    out.append("commit;")
    print(f"{len(alunos)} alunos, {total} notas, {n_listas} listas de preferência", file=sys.stderr)
    return "\n".join(out) + "\n"


if __name__ == "__main__":
    (RAIZ / "supabase" / "02_seed.sql").write_text(seed(), encoding="utf-8")
    print("ok: supabase/02_seed.sql", file=sys.stderr)
    if len(sys.argv) > 1:
        (RAIZ / "supabase" / "03_importar_planilha.sql").write_text(importar(sys.argv[1]), encoding="utf-8")
        print("ok: supabase/03_importar_planilha.sql", file=sys.stderr)
